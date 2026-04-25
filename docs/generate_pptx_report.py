from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Apresentacao_Algoritmo_Genetico_Cancer_de_Mama.pptx"
RESULTS_PATH = ROOT / "artifacts" / "ga_study" / "results.json"
GRAFICOS_DIR = ROOT / "artifacts" / "ga_study" / "plots"


def _pct(value: float) -> str:
    return f"{value * 100:.2f}%".replace(".", ",")


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    text_frame = slide.placeholders[1].text_frame
    text_frame.clear()
    for index, bullet in enumerate(bullets):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(22)


def add_two_column_slide(
    prs: Presentation,
    title: str,
    left_title: str,
    left_bullets: list[str],
    right_title: str,
    right_bullets: list[str],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title

    left_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(4.2), Inches(5.2))
    right_box = slide.shapes.add_textbox(Inches(5.0), Inches(1.4), Inches(4.2), Inches(5.2))

    for box, column_title, bullets in ((left_box, left_title, left_bullets), (right_box, right_title, right_bullets)):
        frame = box.text_frame
        frame.word_wrap = True
        p = frame.paragraphs[0]
        p.text = column_title
        p.font.bold = True
        p.font.size = Pt(24)
        for bullet in bullets:
            item = frame.add_paragraph()
            item.text = bullet
            item.level = 0
            item.font.size = Pt(18)


def add_image_slide(prs: Presentation, title: str, image_path: Path, caption: str | None = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.add_picture(str(image_path), Inches(0.8), Inches(1.3), width=Inches(8.4))
    if caption:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(8.4), Inches(0.5))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = caption
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(14)


def build_presentation() -> Path:
    data = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    models = {item["model_name"]: item for item in data["models"]}
    summary = data["summary_table"]
    knn = models["KNeighborsClassifier"]
    tree = models["DecisionTreeClassifier"]
    logreg = models["LogisticRegression"]
    rf = models["RandomForestClassifier"]
    best_balance = max(
        data["models"],
        key=lambda item: (
            item["optimized_metrics"]["recall"],
            item["optimized_metrics"]["f1_score"],
            item["optimized_metrics"]["specificity"],
        ),
    )
    recall_gain = max(summary, key=lambda row: row["delta_recall"])

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Algoritmo Genetico no Diagnostico de Cancer de Mama em Mulheres",
        "Estudo com 4 modelos baseline e 12 execucoes de otimizacao genetica",
    )

    add_bullet_slide(
        prs,
        "Objetivo do Estudo",
        [
            "Comparar modelos baseline e modelos otimizados por Algoritmo Genetico.",
            "Priorizar recall para reduzir falsos negativos em contexto clinico.",
            "Analisar trade-offs entre recall, especificidade e F1-score.",
        ],
    )

    add_bullet_slide(
        prs,
        "Dataset e Preparacao",
        [
            "Breast Cancer Wisconsin Diagnostic, com 569 amostras.",
            "Split estratificado: 455 em treino e 114 em teste.",
            "Normalizacao com StandardScaler e classe positiva igual a malignidade.",
        ],
    )

    add_bullet_slide(
        prs,
        "Modelos Avaliados",
        [
            "KNeighborsClassifier",
            "DecisionTreeClassifier",
            "LogisticRegression",
            "RandomForestClassifier",
        ],
    )

    add_bullet_slide(
        prs,
        "Configuracao do AG",
        [
            "Fitness = 0,6 x recall + 0,3 x F1-score + 0,1 x especificidade.",
            "Selecao por torneio, crossover, mutacao e elitismo.",
            "3 experimentos por modelo: exploracao, equilibrio e refinamento.",
        ],
    )

    add_two_column_slide(
        prs,
        "Baselines no Teste",
        "Melhores sinais",
        [
            f"LogisticRegression: recall {_pct(logreg['baseline_metrics']['recall'])}",
            f"RandomForest: especificidade {_pct(rf['baseline_metrics']['specificity'])}",
            f"KNN: F1 {_pct(knn['baseline_metrics']['f1_score'])}",
        ],
        "Leitura inicial",
        [
            "LogisticRegression foi o baseline mais forte no conjunto de teste.",
            "RandomForest teve baseline robusto, mas nao foi o melhor em recall.",
            "DecisionTree foi a familia menos equilibrada na partida.",
        ],
    )

    add_two_column_slide(
        prs,
        "Resultado Final por Modelo",
        "Maior ganho de recall",
        [
            f"Modelo: {recall_gain['modelo']}",
            f"Recall baseline: {_pct(recall_gain['recall_baseline'])}",
            f"Recall otimizado: {_pct(recall_gain['recall_otimizado'])}",
            f"Delta: {(recall_gain['delta_recall'] * 100):.2f}".replace(".", ",") + " p.p.",
        ],
        "Melhor equilibrio geral",
        [
            f"Modelo: {best_balance['model_name']}",
            f"Recall: {_pct(best_balance['optimized_metrics']['recall'])}",
            f"Especificidade: {_pct(best_balance['optimized_metrics']['specificity'])}",
            f"F1-score: {_pct(best_balance['optimized_metrics']['f1_score'])}",
        ],
    )

    add_two_column_slide(
        prs,
        "Melhores Experimentos",
        "KNN e Decision Tree",
        [
            f"KNN: {knn['best_experiment_name']}",
            "n_neighbors = 2",
            "weights = distance",
            "p = 1",
            f"Recall final = {_pct(knn['optimized_metrics']['recall'])}",
        ],
        "LogReg e Random Forest",
        [
            f"LogReg: {logreg['best_experiment_name']}",
            "C = 7.5 | max_iter = 600 | solver = liblinear",
            f"RF: {rf['best_experiment_name']}",
            "n_estimators = 400 | max_depth = 30",
            "min_samples_split = 42 | min_samples_leaf = 3",
        ],
    )

    add_two_column_slide(
        prs,
        "Trade-offs Observados",
        "Ganhos",
        [
            "KNN aumentou o recall em 2,38 p.p.",
            "DecisionTree aumentou a especificidade.",
            "LogisticRegression manteve o melhor equilibrio geral.",
        ],
        "Perdas",
        [
            "KNN perdeu especificidade e F1.",
            "DecisionTree perdeu recall.",
            "RandomForest perdeu recall e F1 no teste.",
        ],
    )

    add_image_slide(
        prs,
        "Comparacao Baseline x Otimizado",
        GRAFICOS_DIR / "comparacao_baseline_vs_otimizado.png",
        "Comparacao de recall, especificidade e F1-score para os quatro modelos.",
    )

    add_image_slide(
        prs,
        "Convergencia do KNN",
        GRAFICOS_DIR / "convergencia_KNeighborsClassifier_experimento_1_exploracao.png",
        "O KNN foi a familia com ganho real de recall no conjunto de teste.",
    )

    add_image_slide(
        prs,
        "Convergencia da Logistic Regression",
        GRAFICOS_DIR / "convergencia_LogisticRegression_experimento_1_exploracao.png",
        "A regressao logistica mostrou estabilidade e confirmou uma regiao otima forte.",
    )

    add_image_slide(
        prs,
        "Convergencia do Random Forest",
        GRAFICOS_DIR / "convergencia_RandomForestClassifier_experimento_1_exploracao.png",
        "O ganho em validacao nao se traduziu em melhoria de recall no teste.",
    )

    add_bullet_slide(
        prs,
        "Analise Tecnica",
        [
            "Nem toda melhora de fitness em validacao gera ganho de recall no teste.",
            "KNN foi o unico modelo com ganho real de sensibilidade.",
            "LogisticRegression permaneceu como melhor recomendacao clinica geral.",
        ],
    )

    add_bullet_slide(
        prs,
        "Infraestrutura e Operacao",
        [
            "Interface em Streamlit, API em FastAPI e logging em arquivo.",
            "Persistencia de artefatos em JSON, JSONL e imagens.",
            "Docker, docker-compose e estrutura inicial em Terraform.",
        ],
    )

    add_bullet_slide(
        prs,
        "Conclusao",
        [
            "O estudo comparou 4 modelos baseline e 12 execucoes de AG.",
            "O maior ganho de recall foi do KNN, com 2,38 p.p.",
            "A melhor recomendacao final permaneceu em LogisticRegression.",
            "O projeto ficou coerente com os resultados reais gerados pelo estudo.",
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(OUTPUT))
        return OUTPUT
    except PermissionError:
        fallback = OUTPUT.with_name(f"{OUTPUT.stem}_atualizada{OUTPUT.suffix}")
        prs.save(str(fallback))
        return fallback


if __name__ == "__main__":
    print(build_presentation())
